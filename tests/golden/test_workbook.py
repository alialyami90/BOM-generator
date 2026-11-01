"""The written example workbook (plan M5): every formula, recalculated (Q15), gives the
value the writer cached, and the recalculated totals are the ones the spec states."""

from collections.abc import Callable
from pathlib import Path

import openpyxl
import pytest
from openpyxl.worksheet.worksheet import Worksheet
from pycel import ExcelCompiler

from bomgen.engine.pipeline import build
from bomgen.io.loader import load_substation
from bomgen.sheets.workbook import build_workbook
from bomgen.writer import write_workbook

pytestmark = [
    pytest.mark.golden,
    # pycel 1.0b30 compiles formulas with ast names Python 3.12 deprecates.
    pytest.mark.filterwarnings("ignore:Attribute s is deprecated:DeprecationWarning"),
    pytest.mark.filterwarnings("ignore:ast.Str is deprecated:DeprecationWarning"),
]

SUBSTATION_FILE = (
    Path(__file__).resolve().parents[2] / "examples" / "substations" / "gis01" / "substation.yaml"
)

Evaluate = Callable[[str, str], object]


@pytest.fixture(scope="module")
def path(tmp_path_factory: pytest.TempPathFactory) -> Path:
    loaded = load_substation(SUBSTATION_FILE)
    workbook = build_workbook(loaded, build(loaded))
    return write_workbook(workbook, tmp_path_factory.mktemp("golden") / "gis01.xlsx")


@pytest.fixture(scope="module")
def formulas(path: Path) -> openpyxl.Workbook:
    return openpyxl.load_workbook(path)


# Formula cells in the example workbook; the count changes only if the layout does. The
# Loose Material sheet (ADR 0020) added its own HWIO total; the device nameplate row
# (ADR 0026) adds one SR.NO formula per panel, and the Nameplates sheet one per plate
# after the first (28 plates).
FORMULA_COUNT = 955


def formula_cells(book: openpyxl.Workbook) -> list[tuple[str, str]]:
    """(sheet, cell) of every formula; text such as "=" is not one."""
    return [
        (sheet.title, cell.coordinate)
        for sheet in book.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.data_type == "f"
    ]


@pytest.fixture(scope="module")
def evaluate(path: Path, formulas: openpyxl.Workbook) -> Evaluate:
    """Evaluate a cell from its formula, never from the result the writer cached.

    pycel reads the file's cached results and only calculates a formula that has none,
    so load every formula, then clear and recalculate them all.
    """
    compiler = ExcelCompiler(filename=str(path))
    for sheet, cell in formula_cells(formulas):
        compiler.evaluate(f"{sheet}!{cell}")
    compiler.recalculate()

    def run(sheet: str, cell: str) -> object:
        return compiler.evaluate(f"{sheet}!{cell}")

    return run


def row_of(sheet: Worksheet, column: str, label: str) -> int:
    for cell in sheet[column]:
        if cell.value == label:
            return int(cell.row)
    raise AssertionError(f"{label!r} not found in column {column} of {sheet.title!r}")


def test_sheet_order(formulas: openpyxl.Workbook) -> None:
    # Output requirements (OUT-3) as amended by ADR 0020 (Q36) and ADR 0026: a BOM sheet
    # per cabinet, the Loose Material sheet, the Nameplates sheet, then the three
    # calculation sheets.
    assert formulas.sheetnames == [
        "1", "2", "3", "4", "Loose Material", "Nameplates",
        "Power Consumption", "Heat Dissipation", "Heat Calculations",
    ]  # fmt: skip


def test_header_row_is_53_1_on_the_console_sheet(formulas: openpyxl.Workbook) -> None:
    # Formatting, panel BOM sheets: "Row 26.25 (53.1 on sheet 4)"; the example's sheet 4 is
    # the WS console, read as the console panel (plan Q61).
    # Excel keeps row heights in whole pixels (0.75 pt): XlsxWriter 3.3 writes 53.1 as 53.25.
    heights = [formulas[name].row_dimensions[2].height for name in ("1", "2", "3", "4")]
    assert heights == [26.25, 26.25, 26.25, pytest.approx(53.1, abs=0.375)]


def test_bom_sheet_merges(formulas: openpyxl.Workbook) -> None:
    # Formatting → BOM sheets: the merged title; a designation spanning its kit rows; make
    # and model spanning the kit children that leave them empty.
    merges = {str(m) for m in formulas["1"].merged_cells.ranges}
    assert {"A1:G1", "B4:B7", "F4:F6", "G4:G6"} <= merges


def test_loose_material_sheet(formulas: openpyxl.Workbook) -> None:
    # ADR 0020 (Q36): the loose table leaves the WS console sheet (now A1:G45, was
    # A1:G62) for its own sheet, which is laid out and printed like a panel BOM; the
    # device nameplate row (ADR 0026) adds one row, A1:G46.
    console, loose = formulas["4"], formulas["Loose Material"]
    assert console.print_area == "'4'!$A$1:$G$46"
    assert console.max_row == 46
    assert loose.print_area == "'Loose Material'!$A$1:$G$16"
    assert loose["A1"].value == "GIS Substation\nPSA Panels Loose Material and Auxiliary Equipment"
    assert loose["B2"].value == "ITEM DESIGNATION"
    assert loose["C3"].value == '55" LCD Alarm Annunciator (Touch)  - Wall Mounted'
    assert loose.print_title_rows == "$1:$2"
    assert loose.column_dimensions["H"].hidden
    setup = loose.page_setup
    assert (setup.orientation, setup.paperSize, setup.fitToHeight) == ("portrait", 9, 0)
    fit = loose.sheet_properties.pageSetUpPr
    assert fit is not None
    assert fit.fitToPage
    assert loose.print_options.horizontalCentered


def test_nameplates_sheet(formulas: openpyxl.Workbook) -> None:
    # ADR 0026: every plate's text per panel, laid out and printed like a panel BOM.
    sheet = formulas["Nameplates"]
    assert sheet["A1"].value == "GIS Substation\nDevice Nameplates"
    assert "A1:D1" in {str(m) for m in sheet.merged_cells.ranges}
    assert [sheet.cell(2, c).value for c in range(1, 5)] == [
        "SR.NO",
        "PANEL",
        "PLATE TEXT",
        "DEVICE",
    ]
    assert [sheet.cell(3, c).value for c in range(1, 4)] == [
        1,
        "861-GIS01-PSA-0001",
        "FOPP_CCR_A01",
    ]
    assert sheet["A4"].value == "=A3+1"
    assert [sheet.cell(r, 3).value for r in range(28, 31)] == [
        "L2_SRV/OAS/EAS_A01",  # ADR 0030
        "MONITOR",
        "Keyboard & Mouse",
    ]
    assert sheet.print_area == "'Nameplates'!$A$1:$D$30"
    assert sheet.print_title_rows == "$1:$2"
    setup = sheet.page_setup
    assert (setup.orientation, setup.paperSize, setup.fitToHeight) == ("portrait", 9, 0)
    assert zone(sheet, "oddHeader", "center").splitlines()[1:] == [
        "Bill of Materials for Nameplates",
        "Supplier Doc # VD-EL0009-861GISPSA-001",
    ]


def test_power_consumption_footnote(formulas: openpyxl.Workbook) -> None:
    # Power Consumption sheet, item 5: the printer footnote (output.power_footnotes) two
    # rows below the NOTES, across the printed width, wrapped at its line break, printed.
    sheet = formulas["Power Consumption"]
    notes = row_of(sheet, "B", "NOTES:")
    row = notes + 5
    footnote = sheet.cell(row, 2)
    assert str(footnote.value).startswith("*Printer Power Consumption: 622 watts (Printing)")
    assert str(footnote.value).endswith("Wake on\nLAN);")
    assert footnote.alignment.wrap_text
    assert f"B{row}:L{row}" in {str(m) for m in sheet.merged_cells.ranges}
    assert sheet.row_dimensions[row].height == 30
    assert sheet.print_area == f"'Power Consumption'!$A$1:$L${row}"


def zone(sheet: Worksheet, part: str, side: str) -> str:
    """The text of one header/footer zone, e.g. ``zone(sheet, "oddHeader", "center")``."""
    header_footer = getattr(sheet, part)
    assert header_footer is not None
    text = getattr(header_footer, side).text
    return "" if text is None else str(text)


def test_loose_material_header_and_footer(formulas: openpyxl.Workbook) -> None:
    # ADR 0020: the BOM header with the sheet as its subject, under the host panel's
    # document (the WS console's).
    loose, console = formulas["Loose Material"], formulas["4"]
    assert zone(loose, "oddHeader", "center").splitlines()[1:] == [
        "Bill of Materials for Loose Material",
        "Supplier Doc # VD-EL0009-861GISPSA-001",
    ]
    assert zone(loose, "oddHeader", "right") == zone(console, "oddHeader", "right") == "Rev : 02"
    assert zone(loose, "oddFooter", "left") == "VD-EL0009-861GISPSA-001"
    assert zone(loose, "oddFooter", "center") == "Page &P of &N"


def test_external_equipment_reads_the_loose_material_sheet(
    formulas: openpyxl.Workbook,
) -> None:
    # Power Consumption sheet, item 4: the annunciator row reads its loose-material row,
    # now on the Loose Material sheet (was ='4'!C49).
    sheet = formulas["Power Consumption"]
    row = row_of(sheet, "A", "Power Consumption for External Equipment") + 5
    assert [sheet.cell(row, col).value for col in (5, 6, 7, 8)] == [
        "='Loose Material'!C3", "='Loose Material'!D3",
        "='Loose Material'!E3", "='Loose Material'!F3",
    ]  # fmt: skip


def test_every_formula_recalculates_to_its_cached_value(
    path: Path, formulas: openpyxl.Workbook, evaluate: Evaluate
) -> None:
    # Output requirements (OUT-5): live formulas; the cached results must agree, so the
    # file shows the right numbers before Excel recalculates.
    cached = openpyxl.load_workbook(path, data_only=True)
    cells = formula_cells(formulas)
    wrong = []
    for sheet, cell in cells:
        got, want = evaluate(sheet, cell), cached[sheet][cell].value
        if isinstance(want, int | float) and not isinstance(want, bool):
            same = isinstance(got, int | float) and got == pytest.approx(want)
        else:
            same = got == want
        if not same:
            wrong.append(f"{sheet}!{cell} {formulas[sheet][cell].value}: {got!r} != {want!r}")
    assert len(cells) == FORMULA_COUNT
    assert wrong == []


def test_power_consumption_totals(formulas: openpyxl.Workbook, evaluate: Evaluate) -> None:
    # Power Consumption sheet: 2,177.5 W DC + 1,090.3 W AC = 3,267.8 W.
    sheet = formulas["Power Consumption"]
    dc = row_of(sheet, "C", "125V DC LOAD")
    ac = row_of(sheet, "C", "230V AUX AC LOAD")
    total = row_of(sheet, "E", "TOTAL POWER (W)")
    assert evaluate(sheet.title, f"G{dc}") == pytest.approx(2177.5)
    assert evaluate(sheet.title, f"G{ac}") == pytest.approx(1090.3)
    assert evaluate(sheet.title, f"G{total}") == pytest.approx(3267.8)


def test_required_airflow(formulas: openpyxl.Workbook, evaluate: Evaluate) -> None:
    # Heat Dissipation and Heat Calculations sheets: 391, 391 and 172 m³/h against 716.
    sheet = formulas["Heat Calculations"]
    labels = [c.row for c in sheet["B"] if c.value == "Air Flow Rate in m³/h ="]
    airflow = [evaluate(sheet.title, f"E{row + 1}") for row in labels]
    fans = [evaluate(sheet.title, f"E{row + 3}") for row in labels]
    assert [round(float(str(a))) for a in airflow] == [391, 391, 172]
    assert fans == [716, 716, 716]


def test_line_up_check(formulas: openpyxl.Workbook, evaluate: Evaluate) -> None:
    # Heat Calculations: Σ required = 954.8 m³/h against 3 x 716 = 2,148 m³/h.
    sheet = formulas["Heat Calculations"]
    required = row_of(sheet, "B", "Total required air flow =") + 1
    fans = row_of(sheet, "B", "Total fans capacity =") + 1
    assert evaluate(sheet.title, f"D{required}") == pytest.approx(954.8, abs=0.05)
    assert evaluate(sheet.title, f"D{fans}") == pytest.approx(2148)


def test_console_natural_convection(formulas: openpyxl.Workbook, evaluate: Evaluate) -> None:
    # Heat Calculations: A = 4.848 m² and ΔT = 6.96 °C, under the 10 °C limit.
    sheet = formulas["Heat Calculations"]
    area = row_of(sheet, "B", "Surface area of panel =") + 1
    rise = row_of(sheet, "B", "Heat rise inside the panel =") + 1
    assert evaluate(sheet.title, f"D{area}") == pytest.approx(4.848)
    assert evaluate(sheet.title, f"D{rise}") == pytest.approx(6.96, abs=0.005)
    conclusion = evaluate(sheet.title, f"B{rise + 1}")
    assert isinstance(conclusion, str)
    assert "within the limit" in conclusion
