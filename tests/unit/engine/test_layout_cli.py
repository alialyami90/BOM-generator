"""``bomgen layout`` writes the generated panel files; ``bomgen build`` lays them out again."""

import shutil
from pathlib import Path

import openpyxl
import pytest
import yaml
from typer.testing import CliRunner

from bomgen.cli import app

EXAMPLES = Path(__file__).parents[3] / "examples"
runner = CliRunner()


@pytest.fixture
def substation(tmp_path: Path) -> Path:
    shutil.copytree(EXAMPLES, tmp_path / "examples")
    path = tmp_path / "examples" / "substations" / "gis01" / "substation.yaml"
    # The example feeds the annunciator from the hand-written psa3, which the layout
    # does not write; feed it from psa1 so the layout has no errors.
    path.write_text(path.read_text().replace("fed_from: psa3", "fed_from: psa1"))
    return path


def test_cli_layout_writes_panels_and_prints_fill(substation: Path, tmp_path: Path) -> None:
    out = tmp_path / "generated"
    result = runner.invoke(app, ["layout", "--substation", str(substation), "--out", str(out)])
    assert result.exit_code == 0, result.output
    assert sorted(p.name for p in out.iterdir()) == ["psa1.yaml", "psa2.yaml"]
    assert yaml.safe_load((out / "psa1.yaml").read_text())["generated"] is True
    assert "IED count: 40" in result.stdout
    assert "Panel 1: 29 U / 32 U" in result.stdout


def test_cli_layout_keeps_hand_written_panels(substation: Path) -> None:
    before = (substation.parent / "panels" / "psa1.yaml").read_text()
    result = runner.invoke(app, ["layout", "--substation", str(substation)])
    assert result.exit_code == 1
    assert "hand-written panel file" in result.output
    assert (substation.parent / "panels" / "psa1.yaml").read_text() == before


def test_cli_layout_overwrites_its_own_files(substation: Path) -> None:
    args = ["layout", "--substation", str(substation)]
    assert runner.invoke(app, [*args, "--force"]).exit_code == 0
    assert runner.invoke(app, args).exit_code == 0  # now generated: no --force needed


def test_cli_layout_exits_1_on_layout_errors(substation: Path, tmp_path: Path) -> None:
    text = substation.read_text().replace(
        "  overrides: []", "  overrides: [{pair: gw, panel: [9, 10]}]"
    )
    substation.write_text(text)
    result = runner.invoke(
        app, ["layout", "--substation", str(substation), "--out", str(tmp_path / "g")]
    )
    assert result.exit_code == 1
    assert "error: gw" in result.stdout
    assert "unknown-panel" not in result.stdout  # messages, not kinds


def test_cli_build_after_layout_uses_the_generated_panels(substation: Path, tmp_path: Path) -> None:
    assert runner.invoke(app, ["layout", "--substation", str(substation), "--force"]).exit_code == 0
    out = tmp_path / "out"
    result = runner.invoke(app, ["build", str(substation), "--out", str(out)])
    assert result.exit_code == 0, result.output
    book = openpyxl.load_workbook(next(out.glob("*.xlsx")))
    assert book.sheetnames[:3] == ["1", "2", "4"]  # psa3's devices now sit in psa1 and psa2


def _own_psa1(substation: Path) -> None:
    psa1 = substation.parent / "panels" / "psa1.yaml"
    psa1.write_text(psa1.read_text() + "generated: true\n")


def test_cli_build_no_layout_builds_the_files_as_they_are(substation: Path, tmp_path: Path) -> None:
    _own_psa1(substation)
    out = tmp_path / "out"
    result = runner.invoke(app, ["build", str(substation), "--no-layout", "--out", str(out)])
    assert result.exit_code == 0, result.output
    book = openpyxl.load_workbook(next(out.glob("*.xlsx")))
    assert book.sheetnames[:4] == ["1", "2", "3", "4"]


def test_cli_build_prints_layout_warnings(substation: Path, tmp_path: Path) -> None:
    _own_psa1(substation)  # psa2 and psa3 stay hand-written but listed in panels
    result = runner.invoke(app, ["build", str(substation), "--out", str(tmp_path / "out")])
    assert result.exit_code == 0, result.output
    assert "gis01: warning: psa2:" in result.stderr
    assert "gis01: warning: psa3:" in result.stderr
    assert "IED count" not in result.stderr  # the full report only comes with errors


def _panels_list(substation: Path) -> list[str]:
    return list(yaml.safe_load(substation.read_text())["panels"])


def test_r9_layout_relists_the_panels_it_writes(substation: Path) -> None:
    result = runner.invoke(app, ["layout", "--substation", str(substation), "--force"])
    assert result.exit_code == 0, result.output
    assert _panels_list(substation) == ["panels/psa1.yaml", "panels/psa2.yaml"]
    text = substation.read_text()
    assert "# panel order = sheet order (OUT-3)" in text  # comments survive
    assert "manual_panels:  # never touched by the layout" in text
    # psa3 is hand-written: kept on disk, and the user is told it left the list
    assert (substation.parent / "panels" / "psa3.yaml").exists()
    assert "psa3.yaml: hand-written, no longer listed in panels" in result.stdout


def test_r9_layout_grows_and_shrinks_the_panels_list(substation: Path) -> None:
    args = ["layout", "--substation", str(substation)]
    assert runner.invoke(app, [*args, "--force"]).exit_code == 0
    (substation.parent / "panels" / "psa3.yaml").unlink()  # the old hand-written psa3
    ws = substation.parent / "panels" / "ws.yaml"  # free sheet name '4' for psa4
    ws.write_text(ws.read_text().replace('sheet_name: "4"', 'sheet_name: "WS"'))
    # Push the gateway pair to a new panel pair: the list grows to four files.
    text = substation.read_text()
    substation.write_text(
        text.replace("  overrides: []", "  overrides: [{pair: gw, panel: [3, 4]}]")
    )
    assert runner.invoke(app, args).exit_code == 0
    assert _panels_list(substation)[2:] == ["panels/psa3.yaml", "panels/psa4.yaml"]
    # Drop the override again: the list shrinks and the stale generated files go.
    substation.write_text(substation.read_text().replace("[{pair: gw, panel: [3, 4]}]", "[]"))
    result = runner.invoke(app, args)
    assert result.exit_code == 0, result.output
    assert _panels_list(substation) == ["panels/psa1.yaml", "panels/psa2.yaml"]
    assert not (substation.parent / "panels" / "psa4.yaml").exists()
    assert "removed" in result.stdout


def test_r7_overrides_survive_a_rerun(substation: Path) -> None:
    args = ["layout", "--substation", str(substation)]
    assert runner.invoke(app, [*args, "--force"]).exit_code == 0
    text = substation.read_text().replace(
        "  overrides: []", "  overrides: [{device: gps, panel: 2}]"
    )
    substation.write_text(text)
    assert runner.invoke(app, args).exit_code == 0
    psa2 = substation.parent / "panels" / "psa2.yaml"
    first = psa2.read_text()
    assert "SEL-2488" in first  # the GPS clock moved from psa1
    assert runner.invoke(app, args).exit_code == 0
    assert substation.read_text().count("{device: gps, panel: 2}") == 1
    assert psa2.read_text() == first
